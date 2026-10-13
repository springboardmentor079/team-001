import csv
import io
from datetime import datetime

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.attendance import Attendance
from app.models.expense import Expense
from app.models.inventory import Inventory
from app.models.milestone import ProjectMilestone
from app.models.procurement import Procurement
from app.models.project import Project
from app.models.report import Report
from app.models.resource import Resource
from app.models.user import User
from app.models.worker import Worker


# ------------------------------------------------------------------ summary
def get_summary(db: Session) -> dict:
    total_projects = db.query(func.count(Project.id)).scalar() or 0
    active_projects = (
        db.query(func.count(Project.id)).filter(Project.status == "in_progress").scalar() or 0
    )
    total_budget = db.query(func.coalesce(func.sum(Project.budget), 0)).scalar() or 0
    total_spent = db.query(func.coalesce(func.sum(Expense.amount), 0)).scalar() or 0
    total_workers = db.query(func.count(Worker.id)).scalar() or 0
    present_attendance = (
        db.query(func.count(Attendance.id)).filter(Attendance.status == "present").scalar() or 0
    )
    low_stock_count = (
        db.query(func.count(Inventory.id))
        .filter(Inventory.quantity <= Inventory.reorder_level)
        .scalar()
        or 0
    )
    pending_procurement = (
        db.query(func.count(Procurement.id)).filter(Procurement.status == "requested").scalar() or 0
    )
    projects_by_status = dict(
        db.query(Project.status, func.count(Project.id)).group_by(Project.status).all()
    )

    return {
        "total_projects": total_projects,
        "active_projects": active_projects,
        "total_budget": float(total_budget),
        "total_spent": float(total_spent),
        "total_workers": total_workers,
        "present_attendance_records": present_attendance,
        "low_stock_items": low_stock_count,
        "pending_procurement_requests": pending_procurement,
        "projects_by_status": projects_by_status,
    }


# ---------------------------------------------------------------- analytics
def _count_by(db: Session, column) -> dict:
    return {k: int(v) for k, v in db.query(column, func.count()).group_by(column).all()}


def get_analytics(db: Session) -> dict:
    budget_vs_spent = []
    spent_rows = dict(
        db.query(Expense.project_id, func.sum(Expense.amount)).group_by(Expense.project_id).all()
    )
    for p in db.query(Project).order_by(Project.id).all():
        budget_vs_spent.append({
            "project": p.name,
            "budget": float(p.budget or 0),
            "spent": float(spent_rows.get(p.id, 0) or 0),
        })

    return {
        "projects_by_status": _count_by(db, Project.status),
        "projects_by_category": _count_by(db, Project.category),
        "budget_vs_spent": budget_vs_spent,
        "expenses_by_category": {
            k: float(v) for k, v in db.query(Expense.category, func.sum(Expense.amount)).group_by(Expense.category).all()
        },
        "inventory_by_category": {
            k: float(v) for k, v in db.query(Inventory.category, func.sum(Inventory.quantity)).group_by(Inventory.category).all()
        },
        "procurement_by_status": _count_by(db, Procurement.status),
        "attendance_by_status": _count_by(db, Attendance.status),
        "resources_by_status": _count_by(db, Resource.status),
        "milestones_by_status": _count_by(db, ProjectMilestone.status),
        "workers_by_category": _count_by(db, Worker.category),
    }


# ------------------------------------------------------------------ exports
EXPORT_KINDS = {
    "projects": "Projects report",
    "workers": "Workforce report",
    "attendance": "Attendance report",
    "inventory": "Inventory report",
    "procurement": "Procurement report",
    "expenses": "Expenses report",
    "resources": "Resource utilisation report",
}
# Client users may only export these
CLIENT_KINDS = {"projects", "expenses"}
REPORT_TYPE_BY_KIND = {
    "projects": "progress", "workers": "workforce", "attendance": "workforce",
    "inventory": "resource_utilization", "resources": "resource_utilization",
    "procurement": "procurement", "expenses": "budget",
}


def _num(value) -> float:
    return float(value) if value is not None else 0.0


def build_dataset(db: Session, kind: str) -> dict:
    title = EXPORT_KINDS[kind]
    if kind == "projects":
        columns = ["ID", "Name", "Category", "Status", "Start date", "End date", "Budget"]
        rows = [
            [p.id, p.name, p.category, p.status, p.start_date or "", p.end_date or "", _num(p.budget)]
            for p in db.query(Project).order_by(Project.id).all()
        ]
    elif kind == "workers":
        projects = {p.id: p.name for p in db.query(Project).all()}
        columns = ["ID", "Name", "Category", "Project", "Contact", "Daily wage"]
        rows = [
            [w.id, w.full_name, w.category, projects.get(w.project_id, "-"), w.contact_number or "", _num(w.daily_wage)]
            for w in db.query(Worker).order_by(Worker.id).all()
        ]
    elif kind == "attendance":
        workers = {w.id: w.full_name for w in db.query(Worker).all()}
        projects = {p.id: p.name for p in db.query(Project).all()}
        columns = ["Date", "Worker", "Project", "Status"]
        rows = [
            [a.date, workers.get(a.worker_id, a.worker_id), projects.get(a.project_id, a.project_id), a.status]
            for a in db.query(Attendance).order_by(Attendance.date.desc(), Attendance.id.desc()).all()
        ]
    elif kind == "inventory":
        columns = ["ID", "Material", "Category", "Quantity", "Unit", "Reorder level", "Status"]
        rows = [
            [i.id, i.material_name, i.category, _num(i.quantity), i.unit, _num(i.reorder_level),
             "LOW STOCK" if i.quantity <= i.reorder_level else "OK"]
            for i in db.query(Inventory).order_by(Inventory.id).all()
        ]
    elif kind == "procurement":
        columns = ["ID", "Item", "Category", "Vendor", "Quantity", "Cost", "Status", "Invoice no.", "Paid"]
        rows = [
            [r.id, r.item_name, r.category, r.vendor_name or "", _num(r.quantity), _num(r.cost), r.status,
             r.invoice_number or "", "Yes" if r.invoice_paid else "No"]
            for r in db.query(Procurement).order_by(Procurement.id).all()
        ]
    elif kind == "resources":
        projects = {p.id: p.name for p in db.query(Project).all()}
        columns = ["ID", "Name", "Category", "Assigned project", "Status"]
        rows = [
            [r.id, r.name, r.category, projects.get(r.project_id, "-"), r.status]
            for r in db.query(Resource).order_by(Resource.id).all()
        ]
    else:  # expenses
        projects = {p.id: p.name for p in db.query(Project).all()}
        columns = ["Date", "Project", "Category", "Description", "Amount"]
        rows = [
            [e.expense_date, projects.get(e.project_id, e.project_id), e.category, e.description or "", _num(e.amount)]
            for e in db.query(Expense).order_by(Expense.expense_date.desc(), Expense.id.desc()).all()
        ]
    return {"title": title, "columns": columns, "rows": [[str(c) if hasattr(c, "isoformat") else c for c in r] for r in rows]}


def to_csv(dataset: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(dataset["columns"])
    writer.writerows(dataset["rows"])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(dataset: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = dataset["title"][:31]
    ws.append(dataset["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")
    for row in dataset["rows"]:
        ws.append(row)
    for idx, col in enumerate(ws.columns, start=1):
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(longest + 2, 10), 50)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(dataset: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 8
    cell_style.leading = 10

    def cell(v):
        return Paragraph(str(v).replace("&", "&amp;").replace("<", "&lt;"), cell_style)

    story = [
        Paragraph(f"BuildTrack - {dataset['title']}", styles["Title"]),
        Paragraph(f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')} - {len(dataset['rows'])} record(s)", styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [dataset["columns"]] + (
        [[cell(c) for c in row] for row in dataset["rows"]] or [["No records"] + [""] * (len(dataset["columns"]) - 1)]
    )
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8DDE8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    doc.build(story)
    return out.getvalue()


def log_report(db: Session, kind: str, filename: str, user_id: int) -> Report:
    report = Report(report_type=REPORT_TYPE_BY_KIND[kind], generated_by=user_id, file_path=filename)
    db.add(report)
    db.commit()
    db.refresh(report)
    return report


def report_history(db: Session, limit: int = 20) -> list[dict]:
    rows = (
        db.query(Report, User.full_name)
        .outerjoin(User, User.id == Report.generated_by)
        .order_by(Report.created_at.desc(), Report.id.desc())
        .limit(limit)
        .all()
    )
    return [
        {
            "id": r.id,
            "report_type": r.report_type,
            "file_name": r.file_path,
            "generated_by": name or "-",
            "created_at": r.created_at,
        }
        for r, name in rows
    ]
